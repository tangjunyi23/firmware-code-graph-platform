"""Export Markdown reports to DOCX / PDF.

markdown_to_docx() converts the report Markdown into a .docx with CJK fonts
(宋体 body / 微软雅黑 headings, Table Grid borders). Supported constructs
stay in sync with what pipeline/report.py emits: headings (#..######), pipe
tables, bullet/ordered lists (literal numbers, indentation levels), fenced
code blocks (an unclosed fence degrades to plain paragraphs instead of
swallowing the rest), **bold** (may contain *), `inline code`, [links](url)
rendered as "text (url)", > quotes and --- horizontal rules, plain
paragraphs.

export_report() resolves a report id with the same strict rules as
admin_api._report_path, caches data/reports/<rid>.docx (regenerated when the
source .md is newer), and shells out to LibreOffice
(`soffice --headless --convert-to pdf`) for PDF, cached as <rid>.pdf.
"""

import re
import shutil
import subprocess
import threading
from pathlib import Path

from fastapi import HTTPException

from . import vulnagent_api

# Same strict id shapes as admin_api._RID_RE (no traversal).
_RID_RE = re.compile(r"^job-[a-f0-9]{12}$|^sess-s-[a-z0-9]+-[a-f0-9]{4}$|^pf-[0-9a-f]{8}$")

MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document",
    "pdf": "application/pdf",
}

_EXPORT_LOCK = threading.Lock()
_SOFFICE_TIMEOUT = 600

_TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")  # lazy: bold may contain *
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
_ORDERED_RE = re.compile(r"^\s*(\d+)[.、]\s+(.*)$")
_FENCE_RE = re.compile(r"^\s*```")
_SEP_CELL_RE = re.compile(r"^:?-{2,}:?$")
_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")
_QUOTE_RE = re.compile(r"^>\s?(.*)$")
_HR_RE = re.compile(r"^\s{0,3}(-{3,}|\*{3,}|_{3,})\s*$")


# ---------------------------------------------------------------------------
# markdown -> docx
# ---------------------------------------------------------------------------

def _set_style_font(style, ascii_name, ea_name, size=None, bold=None,
                    black=False):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    style.font.name = ascii_name
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), ea_name)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if black:
        style.font.color.rgb = RGBColor(0, 0, 0)


def _set_run_font(run, ascii_name, ea_name):
    from docx.oxml.ns import qn
    run.font.name = ascii_name
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), ea_name)


def _setup_styles(doc):
    _set_style_font(doc.styles["Normal"], "Times New Roman", "宋体",
                    size=10.5)
    for name, size in (("Heading 1", 18), ("Heading 2", 15),
                       ("Heading 3", 13), ("Heading 4", 12),
                       ("Heading 5", 11), ("Heading 6", 10.5)):
        try:
            _set_style_font(doc.styles[name], "微软雅黑", "微软雅黑",
                            size=size, bold=True, black=True)
        except KeyError:
            pass
    for name in ("List Bullet", "Intense Quote"):
        try:
            _set_style_font(doc.styles[name], "Times New Roman", "宋体",
                            size=10.5)
        except KeyError:
            pass


def _add_runs(paragraph, text, force_bold=False):
    """Add runs to a paragraph, honoring **bold**, `inline code` and links.

    Markdown links [text](url) degrade to "text (url)" so the target stays
    visible in the exported document."""
    text = _LINK_RE.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    for tok in _TOKEN_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
        elif (tok.startswith("`") and tok.endswith("`") and len(tok) > 2
                and tok[1:-1].strip("`")):  # "```" alone is not inline code
            run = paragraph.add_run(tok[1:-1])
            _set_run_font(run, "Consolas", "宋体")
        else:
            run = paragraph.add_run(tok)
            if force_bold:
                run.bold = True


def _add_horizontal_rule(doc):
    """A ---/*** line becomes a paragraph whose bottom border is the rule."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    paragraph = doc.add_paragraph()
    p_pr = paragraph._element.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "808080")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_quote(doc, text):
    """A > line becomes an Intense Quote paragraph (indent fallback)."""
    from docx.shared import Pt
    try:
        paragraph = doc.add_paragraph(style="Intense Quote")
    except KeyError:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
    _add_runs(paragraph, text)


def _add_code_block(doc, block_lines):
    from docx.shared import Pt
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(18)
    for i, line in enumerate(block_lines):
        run = paragraph.add_run(line)
        _set_run_font(run, "Consolas", "宋体")
        run.font.size = Pt(9.5)
        if i != len(block_lines) - 1:
            run.add_break()


def _is_sep_row(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(_SEP_CELL_RE.match(c) for c in cells)


def _split_row(line: str) -> list:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _add_table(doc, rows):
    from docx.oxml.ns import qn
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(cols):
            text = row[c] if c < len(row) else ""
            cell = table.cell(r, c)
            paragraph = cell.paragraphs[0]
            _add_runs(paragraph, text, force_bold=(r == 0))
            for run in paragraph.runs:
                rpr = run._element.get_or_add_rPr()
                rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def markdown_to_docx(md_text: str, out_path) -> "Path":
    """Convert report Markdown to a .docx file at out_path."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    _setup_styles(doc)
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if _FENCE_RE.match(line):
            # look for the closing fence first: an unclosed ``` must NOT
            # swallow the rest of the document into a code block — treat
            # the fence line itself as a plain paragraph instead
            j = i + 1
            while j < len(lines) and not _FENCE_RE.match(lines[j]):
                j += 1
            if j < len(lines):
                _add_code_block(doc, lines[i + 1:j])
                i = j + 1
                continue
            # unclosed: fall through, the ``` line renders as a paragraph
        if line.lstrip().startswith("|") and i + 1 < len(lines) \
                and _is_sep_row(lines[i + 1]):
            rows = [_split_row(line)]
            i += 2  # header + separator
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            _add_table(doc, rows)
            continue
        if _HR_RE.match(line):
            _add_horizontal_rule(doc)
            i += 1
            continue
        m = _HEADING_RE.match(line)
        if m:
            paragraph = doc.add_heading("", level=len(m.group(1)))
            _add_runs(paragraph, m.group(2).strip())
            i += 1
            continue
        m = _QUOTE_RE.match(line)
        if m:
            _add_quote(doc, m.group(1).strip())
            i += 1
            continue
        m = _BULLET_RE.match(line)
        if m:
            paragraph = doc.add_paragraph(style="List Bullet")
            _add_runs(paragraph, m.group(1).strip())
            i += 1
            continue
        m = _ORDERED_RE.match(line)
        if m:
            # literal number keeps每个 ordered list 的编号与 md 原文一致;
            # 缩进层级按前导空白换算 (2 空格或 1 tab 一级)
            expanded = line.expandtabs(2)
            level = (len(expanded) - len(expanded.lstrip())) // 2
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18 * (1 + level))
            _add_runs(paragraph, f"{m.group(1)}. {m.group(2).strip()}")
            i += 1
            continue
        if line.strip():
            paragraph = doc.add_paragraph()
            _add_runs(paragraph, line.strip())
        i += 1
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# ---------------------------------------------------------------------------
# markdown -> pdf（无 LibreOffice 时的纯 Python 兜底，reportlab + STSong CID 字体）
# ---------------------------------------------------------------------------

_PDF_FONTS_READY = False


def _ensure_pdf_fonts():
    global _PDF_FONTS_READY
    if _PDF_FONTS_READY:
        return
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    except ImportError as exc:  # pragma: no cover - 部署缺依赖时给出可操作的提示
        raise HTTPException(
            status_code=503,
            detail="服务器未安装 reportlab,无法导出 PDF;"
                   "请在 venv 内执行 pip install reportlab") from exc
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    _PDF_FONTS_READY = True


def _pdf_inline(text: str) -> str:
    """md 行内标记 -> platypus 迷你标记：链接转明文、**粗体** 去标记、`code` 等宽。"""
    from xml.sax.saxutils import escape

    text = _LINK_RE.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    out = []
    for tok in _TOKEN_RE.split(text):
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            out.append(escape(tok[2:-2]))
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            out.append("<font face='Courier' size='9'>%s</font>"
                       % escape(tok[1:-1]))
        else:
            out.append(escape(tok))
    return "".join(out)


def markdown_to_pdf(md_text: str, out_path) -> "Path":
    """report Markdown -> PDF（reportlab）。支持与 DOCX 相同的构造：
    标题/表格/列表/围栏代码/引用/分隔线/段落。"""
    _ensure_pdf_fonts()
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    Table, TableStyle, HRFlowable,
                                    Preformatted)

    cn = "STSong-Light"
    body = ParagraphStyle("body", fontName=cn, fontSize=10.5, leading=16.5,
                          textColor=colors.HexColor("#1F2937"), spaceAfter=6)
    quote = ParagraphStyle("quote", parent=body, leftIndent=14,
                           textColor=colors.HexColor("#6B7280"))
    code = ParagraphStyle("code", fontName=cn, fontSize=8.5, leading=12.5,
                          backColor=colors.HexColor("#F4F6F8"),
                          borderColor=colors.HexColor("#E5E7EB"),
                          borderWidth=0.5, borderPadding=6,
                          leftIndent=6, rightIndent=6, spaceBefore=4,
                          spaceAfter=8)
    hstyles = []
    for lvl, (size, color) in enumerate([
            (18, "#0F766E"), (15, "#0F766E"), (13, "#155E75"),
            (12, "#374151"), (11, "#374151"), (10.5, "#374151")], start=1):
        hstyles.append(ParagraphStyle(
            f"h{lvl}", fontName=cn, fontSize=size,
            leading=size * 1.4, textColor=colors.HexColor(color),
            spaceBefore=12 if lvl <= 2 else 8, spaceAfter=5))

    story = []
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if _FENCE_RE.match(line):
            j = i + 1
            while j < len(lines) and not _FENCE_RE.match(lines[j]):
                j += 1
            if j < len(lines):
                block = "\n".join(lines[i + 1:j])
                story.append(Preformatted(block, code))
                i = j + 1
                continue
        if line.lstrip().startswith("|") and i + 1 < len(lines) \
                and _is_sep_row(lines[i + 1]):
            rows = [_split_row(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            ncols = max(len(r) for r in rows)
            data = [[Paragraph(_pdf_inline(c), body) for c in
                     r + [""] * (ncols - len(r))] for r in rows]
            table = Table(data, colWidths=[(A4[0] - 80) / ncols] * ncols,
                          repeatRows=1)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5,
                 colors.HexColor("#D1D5DB")),
                ("BACKGROUND", (0, 0), (-1, 0),
                 colors.HexColor("#EEF2F7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(table)
            story.append(Spacer(1, 6))
            continue
        if _HR_RE.match(line):
            story.append(HRFlowable(width="100%", thickness=0.6,
                                    color=colors.HexColor("#D1D5DB"),
                                    spaceBefore=8, spaceAfter=8))
            i += 1
            continue
        m = _HEADING_RE.match(line)
        if m:
            story.append(Paragraph(_pdf_inline(m.group(2).strip()),
                                   hstyles[len(m.group(1)) - 1]))
            i += 1
            continue
        m = _QUOTE_RE.match(line)
        if m:
            story.append(Paragraph(_pdf_inline(m.group(1).strip()), quote))
            i += 1
            continue
        m = _BULLET_RE.match(line)
        if m:
            story.append(Paragraph("• " + _pdf_inline(m.group(1).strip()), body))
            i += 1
            continue
        m = _ORDERED_RE.match(line)
        if m:
            expanded = line.expandtabs(2)
            level = (len(expanded) - len(expanded.lstrip())) // 2
            st = ParagraphStyle(f"ol{level}", parent=body,
                                leftIndent=14 * (1 + level))
            story.append(Paragraph(f"{m.group(1)}. "
                                   + _pdf_inline(m.group(2).strip()), st))
            i += 1
            continue
        if line.strip():
            story.append(Paragraph(_pdf_inline(line.strip()), body))
        i += 1

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                            leftMargin=40, rightMargin=40,
                            topMargin=44, bottomMargin=44,
                            title="fwgraph report")
    doc.build(story)
    return out_path


# ---------------------------------------------------------------------------
# report id resolution + cached export
# ---------------------------------------------------------------------------

def report_md_path(rid: str, data_dir) -> Path:
    """Resolve a report id to its Markdown source (strict, no traversal)."""
    if not _RID_RE.match(rid):
        raise HTTPException(status_code=400, detail="bad report id format")
    if rid.startswith("job-") or rid.startswith("pf-"):
        path = Path(data_dir) / "reports" / f"{rid}.md"
    else:
        path = (vulnagent_api.VULNAGENT_HOME / "sessions"
                / rid[len("sess-"):] / "report.md")
    if not path.is_file():
        raise HTTPException(status_code=404, detail="report not found")
    return path


def _soffice_bin() -> str:
    return shutil.which("soffice") or shutil.which("libreoffice") or ""


def export_report(rid: str, fmt: str, data_dir) -> Path:
    """Export report `rid` as `fmt` (docx|pdf); returns the cached file."""
    if fmt not in MEDIA_TYPES:
        raise HTTPException(status_code=400, detail="fmt 仅支持 docx|pdf")
    data_dir = Path(data_dir)
    md = report_md_path(rid, data_dir)
    reports_dir = data_dir / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)
    with _EXPORT_LOCK:
        docx_path = reports_dir / f"{rid}.docx"
        if (not docx_path.is_file()
                or docx_path.stat().st_mtime < md.stat().st_mtime):
            markdown_to_docx(md.read_text(encoding="utf-8", errors="replace"),
                             docx_path)
        if fmt == "docx":
            return docx_path
        pdf_path = reports_dir / f"{rid}.pdf"
        if (not pdf_path.is_file()
                or pdf_path.stat().st_mtime < docx_path.stat().st_mtime):
            soffice = _soffice_bin()
            if not soffice:
                # 无 LibreOffice（如容器/最小化部署）：纯 Python 渲染兜底
                markdown_to_pdf(md.read_text(encoding="utf-8",
                                             errors="replace"), pdf_path)
                return pdf_path
            try:
                subprocess.run(
                    [soffice, "--headless", "--convert-to", "pdf",
                     "--outdir", str(reports_dir), str(docx_path)],
                    check=True, capture_output=True,
                    timeout=_SOFFICE_TIMEOUT)
            except subprocess.TimeoutExpired:
                raise HTTPException(status_code=503,
                                    detail="PDF 转换超时,请稍后重试")
            except (subprocess.CalledProcessError, OSError) as exc:
                raise HTTPException(
                    status_code=500,
                    detail=f"PDF 转换失败:{str(exc)[:200]}")
            if not pdf_path.is_file():
                raise HTTPException(status_code=500,
                                    detail="PDF 转换失败:未生成输出文件")
        return pdf_path
