"""Unit tests for DOCX/PDF report export (report_export.py + endpoint).

Same isolation pattern as test_admin_api.py: temp FWGRAPH_DATA +
VULNAGENT_HOME + legacy ORCH_TOKEN. The PDF test is skipped when
LibreOffice (soffice) is not installed on the host.
"""

import io
import json
import shutil
import zipfile

import pytest
from fastapi.testclient import TestClient

from orchestrator.app import main, vulnagent_api

docx = pytest.importorskip("docx", reason="python-docx not installed")  # noqa: F401

LEGACY = "orch-test-token"
JOB = "abcdef012345"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("FWGRAPH_DATA", str(tmp_path))
    monkeypatch.setenv("ORCH_TOKEN", LEGACY)
    vulnagent_home = tmp_path / "vulnagent"
    (vulnagent_home / "sessions").mkdir(parents=True)
    (vulnagent_home / "findings").mkdir(parents=True)
    monkeypatch.setenv("VULNAGENT_HOME", str(vulnagent_home))
    monkeypatch.setattr(vulnagent_api, "VULNAGENT_HOME", vulnagent_home)
    yield TestClient(main.app)
    with main._jobs_lock:
        main._jobs.pop(JOB, None)


def _legacy():
    return {"Authorization": f"Bearer {LEGACY}"}


def _make_job(tmp_path):
    job = {"job_id": JOB, "firmware": "mx12.bin", "status": "routed",
           "error": None, "created_at": "t0", "updated_at": "t0",
           "size_bytes": 1}
    with main._jobs_lock:
        main._jobs[JOB] = job
    job_dir = tmp_path / "firmware" / JOB
    job_dir.mkdir(parents=True)
    (job_dir / "job.json").write_text(json.dumps(job), encoding="utf-8")


def _gen_report(client, tmp_path):
    _make_job(tmp_path)
    resp = client.post(f"/jobs/{JOB}/report", headers=_legacy())
    assert resp.status_code == 200


class TestReportExport:
    def test_export_docx_valid_zip_with_chinese(self, client, tmp_path):
        _gen_report(client, tmp_path)
        resp = client.get(f"/reports/job-{JOB}/export?fmt=docx",
                          headers=_legacy())
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument")
        assert resp.content[:2] == b"PK"
        with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        assert "漏洞报告" in xml
        assert "mx12.bin" in xml

    def test_export_docx_cached_regenerated(self, client, tmp_path):
        _gen_report(client, tmp_path)
        url = f"/reports/job-{JOB}/export?fmt=docx"
        first = client.get(url, headers=_legacy())
        assert first.status_code == 200
        assert client.get(url, headers=_legacy()).status_code == 200

    def test_export_pdf(self, client, tmp_path):
        if not (shutil.which("soffice") or shutil.which("libreoffice")):
            pytest.skip("LibreOffice (soffice) not installed")
        _gen_report(client, tmp_path)
        resp = client.get(f"/reports/job-{JOB}/export?fmt=pdf",
                          headers=_legacy())
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/pdf")
        assert resp.content[:5] == b"%PDF-"

    def test_export_bad_fmt_400(self, client, tmp_path):
        _gen_report(client, tmp_path)
        resp = client.get(f"/reports/job-{JOB}/export?fmt=html",
                          headers=_legacy())
        assert resp.status_code == 400

    def test_export_traversal_rejected(self, client):
        resp = client.get("/reports/job-../../etc/export?fmt=docx",
                          headers=_legacy())
        assert resp.status_code in (400, 404)

    def test_export_missing_report_404(self, client):
        resp = client.get("/reports/job-000000000000/export?fmt=docx",
                          headers=_legacy())
        assert resp.status_code == 404

    def test_export_requires_auth(self, client, tmp_path):
        _gen_report(client, tmp_path)
        resp = client.get(f"/reports/job-{JOB}/export?fmt=docx")
        assert resp.status_code == 401


class TestMarkdownConstructs:
    """markdown_to_docx construct coverage: H5/H6, links, quotes, rules,
    bold-with-stars, unclosed fences, ordered-list numbering/indentation."""

    def _convert(self, tmp_path, md):
        from docx import Document
        from orchestrator.app import report_export
        out = tmp_path / "out.docx"
        report_export.markdown_to_docx(md, out)
        return Document(str(out))

    def test_heading_5_and_6(self, tmp_path):
        doc = self._convert(tmp_path, "##### 五级标题\n\n###### 六级标题")
        styles = [p.style.name for p in doc.paragraphs]
        assert styles == ["Heading 5", "Heading 6"]
        assert [p.text for p in doc.paragraphs] == ["五级标题", "六级标题"]

    def test_link_renders_text_and_url(self, tmp_path):
        doc = self._convert(tmp_path, "见 [详情页](http://a.b/c) 说明")
        assert doc.paragraphs[0].text == "见 详情页 (http://a.b/c) 说明"

    def test_quote_style(self, tmp_path):
        doc = self._convert(tmp_path, "> 这是一段引用")
        assert doc.paragraphs[0].style.name in ("Intense Quote", "Quote")
        assert doc.paragraphs[0].text == "这是一段引用"

    def test_horizontal_rule(self, tmp_path):
        doc = self._convert(tmp_path, "上文\n\n---\n\n下文")
        texts = [p.text for p in doc.paragraphs]
        assert "---" not in texts
        assert texts[0] == "上文" and texts[-1] == "下文"
        assert "w:pBdr" in doc.element.xml  # the rule is a bottom border

    def test_bold_containing_star(self, tmp_path):
        doc = self._convert(tmp_path, "前缀 **加 * 粗** 后缀")
        runs = doc.paragraphs[0].runs
        bold = [r.text for r in runs if r.bold]
        assert bold == ["加 * 粗"]
        assert "".join(r.text for r in runs) == "前缀 加 * 粗 后缀"

    def test_unclosed_fence_does_not_swallow(self, tmp_path):
        doc = self._convert(tmp_path, "```\n代码行\n\n后续段落")
        texts = [p.text for p in doc.paragraphs]
        assert texts == ["```", "代码行", "后续段落"]
        # nothing rendered as a code block (Consolas runs)
        for p in doc.paragraphs:
            for run in p.runs:
                assert run.font.name != "Consolas"

    def test_closed_fence_still_code(self, tmp_path):
        doc = self._convert(tmp_path, "```\nchar *p;\n```")
        assert len(doc.paragraphs) == 1
        assert doc.paragraphs[0].runs[0].font.name == "Consolas"
        assert doc.paragraphs[0].text == "char *p;"

    def test_ordered_list_numbering_and_indent(self, tmp_path):
        from docx.shared import Pt
        doc = self._convert(tmp_path, "3. 顶层\n\n  7. 嵌套")
        first, second = doc.paragraphs
        assert first.text == "3. 顶层"  # literal number preserved
        assert first.paragraph_format.left_indent == Pt(18)
        assert second.text == "7. 嵌套"
        assert second.paragraph_format.left_indent == Pt(36)
